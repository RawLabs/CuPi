"""Print and editable document exports from the same Qt Markdown model."""
from pathlib import Path

from PySide6.QtCore import QMarginsF, QUrl
from PySide6.QtGui import (
    QFont, QImage, QPageLayout, QPageSize, QTextCursor, QTextDocument, QTextTable,
)
from PySide6.QtPrintSupport import QPrinter
from ai.safe_markdown import SafeMarkdownDocument, sanitize_markdown
import re


def markdown_document(markdown: str, directory: Path) -> QTextDocument:
    document = SafeMarkdownDocument()
    allowed = []
    for image_path in (directory / "images").glob("step_*.png"):
        if re.fullmatch(r"step_\d+\.png", image_path.name) and not image_path.is_symlink():
            if image_path.resolve().parent == (directory / "images").resolve():
                allowed.append(f"./images/{image_path.name}")
    markdown = sanitize_markdown(markdown, allowed_images=allowed)
    document.setDefaultFont(QFont("Sans Serif", 11))
    document.setBaseUrl(QUrl.fromLocalFile(str(directory) + "/"))
    document.setMarkdown(
        markdown,
        QTextDocument.MarkdownFeature.MarkdownDialectGitHub
        | QTextDocument.MarkdownFeature.MarkdownNoHTML,
    )
    return document


def export_pdf(markdown: str, directory: Path, path: Path) -> None:
    document = markdown_document(markdown, directory)
    # Limit evidence to the printable area, keeping portrait/tall screenshots
    # on a page while preserving the original PNGs in the bundle.
    block = document.begin()
    while block.isValid():
        iterator = block.begin()
        while not iterator.atEnd():
            fragment = iterator.fragment()
            if fragment.isValid() and fragment.charFormat().isImageFormat():
                fmt = fragment.charFormat().toImageFormat()
                image = QImage(str(directory / fmt.name()))
                if not image.isNull():
                    document.addResource(QTextDocument.ResourceType.ImageResource, QUrl(fmt.name()), image)
                    document.addResource(QTextDocument.ResourceType.ImageResource, QUrl.fromLocalFile(str((directory / fmt.name()).resolve())), image)
                    scale = min(1.0, 650 / image.width(), 850 / image.height())
                    fmt.setWidth(image.width() * scale)
                    fmt.setHeight(image.height() * scale)
                    cursor = QTextCursor(document)
                    cursor.setPosition(fragment.position())
                    cursor.setPosition(fragment.position() + fragment.length(), QTextCursor.MoveMode.KeepAnchor)
                    cursor.setCharFormat(fmt)
            iterator += 1
        block = block.next()
    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    printer.setOutputFileName(str(path))
    printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    printer.setPageMargins(QMarginsF(18, 18, 18, 18), QPageLayout.Unit.Millimeter)
    document.print_(printer)
    if not path.exists() or path.stat().st_size == 0:
        raise OSError(f"Could not write PDF to {path}")


def export_docx(markdown: str, directory: Path, path: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    source = markdown_document(markdown, directory)
    output = Document()
    section = output.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"):
        style = output.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    output.styles["Normal"].font.size = Pt(11)
    output.styles["Normal"].paragraph_format.space_after = Pt(6)

    def add_block(block, target):
        heading = block.blockFormat().headingLevel()
        paragraph = target.add_paragraph()
        if heading:
            paragraph.style = "Title" if heading == 1 and block.position() == 0 else f"Heading {min(heading, 6)}"
        elif block.textList():
            text_list = block.textList()
            ordered = text_list.format().style().value <= -4
            paragraph.add_run(f"{text_list.itemNumber(block) + 1}. " if ordered else "• ")
            paragraph.paragraph_format.left_indent = Inches(0.2 * text_list.format().indent())
        iterator = block.begin()
        while not iterator.atEnd():
            fragment = iterator.fragment()
            if fragment.isValid():
                fmt = fragment.charFormat()
                if fmt.isImageFormat():
                    image_path = directory / fmt.toImageFormat().name()
                    image = QImage(str(image_path))
                    if not image.isNull():
                        width = min(6.7, image.width() / 96, 8 * image.width() / image.height())
                        paragraph.add_run().add_picture(str(image_path), width=Inches(width))
                else:
                    run = paragraph.add_run(fragment.text().replace("\u2028", "\n"))
                    run.bold = fmt.fontWeight() >= QFont.Weight.Bold.value
                    run.italic = fmt.fontItalic()
                    run.underline = fmt.fontUnderline()
                    run.font.strike = fmt.fontStrikeOut()
                    if fmt.fontFixedPitch() or block.blockFormat().nonBreakableLines():
                        run.font.name = "Consolas"
                        run.font.size = Pt(10)
                    if fmt.isAnchor() and fmt.anchorHref():
                        # Native hyperlinks remain clickable in Word.
                        from docx.opc.constants import RELATIONSHIP_TYPE
                        relationship = paragraph.part.relate_to(fmt.anchorHref(), RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                        link = OxmlElement("w:hyperlink")
                        link.set(qn("r:id"), relationship)
                        link.append(run._r)
                        paragraph._p.append(link)
            iterator += 1

    def add_frame(frame, target):
        iterator = frame.begin()
        while not iterator.atEnd():
            child = iterator.currentFrame()
            if isinstance(child, QTextTable):
                table = target.add_table(rows=child.rows(), cols=child.columns())
                table.style = "Table Grid"
                for row in range(child.rows()):
                    for col in range(child.columns()):
                        cell = table.cell(row, col)
                        source_cell = child.cellAt(row, col)
                        block = source_cell.firstCursorPosition().block()
                        last_position = source_cell.lastCursorPosition().position()
                        while block.isValid() and block.position() <= last_position:
                            add_block(block, cell)
                            block = block.next()
                        # python-docx creates a mandatory empty paragraph.
                        if len(cell.paragraphs) > 1:
                            cell._tc.remove(cell.paragraphs[0]._p)
                        if row == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.bold = True
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
            elif child:
                add_frame(child, target)
            else:
                block = iterator.currentBlock()
                if block.isValid():
                    add_block(block, target)
            iterator += 1

    add_frame(source.rootFrame(), output)
    output.save(path)
